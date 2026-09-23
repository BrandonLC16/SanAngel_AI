"""Bounded, read-only parsing of the versioned branch price workbook."""

import re
import warnings
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
from pydantic import ValidationError

from backend.app.schemas.price import PriceData, normalize_price_unit
from backend.app.schemas.price_import_template import (
    PRICE_IMPORT_BRANCH_CELL,
    PRICE_IMPORT_BRANCH_LABEL_CELL,
    PRICE_IMPORT_COLUMNS,
    PRICE_IMPORT_FIRST_DATA_ROW,
    PRICE_IMPORT_HEADER_ROW,
    PRICE_IMPORT_SCHEMA_VERSION,
    PRICE_IMPORT_SHEET,
    PRICE_IMPORT_VERSION_CELL,
    PRICE_IMPORT_VERSION_LABEL_CELL,
)
from backend.app.services.branch_scope import BranchScope

MAX_PRICE_IMPORT_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 10 * 1024 * 1024
MAX_PACKAGE_MEMBERS = 128
MAX_PRICE_IMPORT_ROWS = 1000
MAX_PRODUCT_ID = 2**63 - 1
LAST_DATA_ROW = PRICE_IMPORT_FIRST_DATA_ROW + MAX_PRICE_IMPORT_ROWS - 1
CELL_REFERENCE = re.compile(r"([A-Z]+)([1-9][0-9]*)")


@dataclass(frozen=True, slots=True)
class PriceImportIssue:
    row: int | None
    field: str | None
    code: str


@dataclass(frozen=True, slots=True)
class PriceImportRow:
    source_row: int
    product_id: int
    product_name: str
    unit: str
    price_mxn: Decimal
    verified_on: date


@dataclass(frozen=True, slots=True)
class PriceImportResult:
    rows: tuple[PriceImportRow, ...]
    issues: tuple[PriceImportIssue, ...]

    @property
    def is_valid(self) -> bool:
        return not self.issues


def _file_issue(code: str) -> PriceImportResult:
    return PriceImportResult((), (PriceImportIssue(None, None, code),))


def _safe_package(package: ZipFile) -> bool:
    members = package.infolist()
    if not members or len(members) > MAX_PACKAGE_MEMBERS:
        return False
    names: set[str] = set()
    total_size = 0
    for member in members:
        name = member.filename
        normalized = name.lower()
        if (
            not name
            or name.startswith("/")
            or "\\" in name
            or ":" in name
            or any(part in {"", ".", ".."} for part in name.rstrip("/").split("/"))
            or normalized in names
            or member.flag_bits & 1
            or "vba" in normalized
            or "macro" in normalized
            or "externallink" in normalized
            or "activex" in normalized
            or "embeddings" in normalized
        ):
            return False
        names.add(normalized)
        total_size += member.file_size
        if total_size > MAX_UNCOMPRESSED_BYTES:
            return False

    if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
        return False
    if b"macroenabled" in package.read("[Content_Types].xml").lower():
        return False

    for member in members:
        name = member.filename
        if name.endswith(".rels"):
            root = ElementTree.fromstring(package.read(name))
            if any(element.attrib.get("TargetMode") == "External" for element in root):
                return False
        if name.startswith("xl/worksheets/") and name.endswith(".xml"):
            root = ElementTree.fromstring(package.read(name))
            if any(element.tag.rsplit("}", 1)[-1] == "f" for element in root.iter()):
                return False
            for element in root.iter():
                if element.tag.rsplit("}", 1)[-1] != "c":
                    continue
                match = CELL_REFERENCE.fullmatch(element.attrib.get("r", ""))
                if (
                    match is None
                    or int(match[2]) > LAST_DATA_ROW
                    or match[1] not in {"A", "B", "C", "D", "E", "F"}
                ):
                    return False
    return package.testzip() is None


def _parse_row(
    row_number: int, cells: tuple[object, ...]
) -> tuple[PriceImportRow | None, list[PriceImportIssue]]:
    values = [cell.value for cell in cells]
    issues: list[PriceImportIssue] = []

    def issue(field: str, code: str) -> None:
        issues.append(PriceImportIssue(row_number, field, code))

    if values[5] is not None:
        issue("extra_column", "unexpected_value")

    product_id = values[0]
    if type(product_id) is not int or not 1 <= product_id <= MAX_PRODUCT_ID:
        issue("product_id", "invalid_positive_integer")

    product_name = values[1]
    if (
        not isinstance(product_name, str)
        or not 1 <= len(product_name) <= 120
        or product_name != product_name.strip()
        or not product_name.isprintable()
        or product_name.startswith(("=", "+", "-", "@"))
    ):
        issue("product_name", "invalid_text")

    unit = values[2]
    if not isinstance(unit, str) or unit != unit.strip().lower():
        issue("unit", "invalid_code")
    else:
        try:
            normalize_price_unit(unit)
        except ValueError:
            issue("unit", "invalid_code")

    price = values[3]
    amount: Decimal | None = None
    if type(price) not in {int, float}:
        issue("price_mxn", "invalid_number")
    else:
        try:
            amount = Decimal(str(price))
            PriceData(amount=amount, unit="kg")
        except (ValueError, ValidationError):
            issue("price_mxn", "out_of_range")

    verified_on = values[4]
    if isinstance(verified_on, datetime):
        verified_date = verified_on.date() if verified_on.time() == time.min else None
    elif isinstance(verified_on, date):
        verified_date = verified_on
    else:
        verified_date = None
    if verified_date is None or verified_date > date.today():
        issue("verified_on", "invalid_date")

    if issues:
        return None, issues
    return (
        PriceImportRow(row_number, product_id, product_name, unit, amount, verified_date),
        issues,
    )


def parse_price_import(
    content: bytes, *, filename: str, branch_scope: BranchScope
) -> PriceImportResult:
    """Parse bytes only; a rejected file yields no rows and never touches the database."""

    if (
        not isinstance(filename, str)
        or not filename.lower().endswith(".xlsx")
        or filename.startswith(".")
        or ".." in filename
        or any(character in filename for character in ("/", "\\", ":", "\x00"))
    ):
        return _file_issue("invalid_filename")
    if not isinstance(content, bytes) or not 0 < len(content) <= MAX_PRICE_IMPORT_BYTES:
        return _file_issue("invalid_size")

    try:
        with ZipFile(BytesIO(content)) as package:
            if not _safe_package(package):
                return _file_issue("unsafe_package")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=False, keep_links=False, keep_vba=False
            )
        try:
            if workbook.sheetnames != [PRICE_IMPORT_SHEET]:
                return _file_issue("invalid_sheets")
            sheet = workbook[PRICE_IMPORT_SHEET]
            # Read-only mode otherwise trusts a workbook-supplied dimension that can hide rows.
            sheet.reset_dimensions()
            if (
                sheet[PRICE_IMPORT_VERSION_LABEL_CELL].value != "schema_version"
                or type(sheet[PRICE_IMPORT_VERSION_CELL].value) is not int
                or sheet[PRICE_IMPORT_VERSION_CELL].value != PRICE_IMPORT_SCHEMA_VERSION
                or sheet[PRICE_IMPORT_BRANCH_LABEL_CELL].value != "branch_code"
            ):
                return _file_issue("invalid_metadata")
            if sheet[PRICE_IMPORT_BRANCH_CELL].value != branch_scope.branch_code:
                return _file_issue("branch_mismatch")
            headers = tuple(sheet.cell(PRICE_IMPORT_HEADER_ROW, col).value for col in range(1, 6))
            if (
                headers != PRICE_IMPORT_COLUMNS
                or sheet.cell(PRICE_IMPORT_HEADER_ROW, 6).value is not None
            ):
                return _file_issue("invalid_headers")
            rows: list[PriceImportRow] = []
            issues: list[PriceImportIssue] = []
            seen: set[tuple[int, str]] = set()
            for row_number, cells in enumerate(
                sheet.iter_rows(
                    min_row=PRICE_IMPORT_FIRST_DATA_ROW, max_row=LAST_DATA_ROW, max_col=6
                ),
                start=PRICE_IMPORT_FIRST_DATA_ROW,
            ):
                if all(cell.value is None for cell in cells):
                    continue
                parsed, row_issues = _parse_row(row_number, cells)
                issues.extend(row_issues)
                if parsed is not None:
                    key = (parsed.product_id, parsed.unit)
                    if key in seen:
                        issues.append(PriceImportIssue(row_number, "product_id/unit", "duplicate"))
                    else:
                        seen.add(key)
                        rows.append(parsed)
            if not rows and not issues:
                return _file_issue("empty_data")
            return PriceImportResult(tuple(rows) if not issues else (), tuple(issues))
        finally:
            workbook.close()
    except Exception:
        # Untrusted package/parser failures never expose paths or source values to callers.
        return _file_issue("invalid_workbook")
