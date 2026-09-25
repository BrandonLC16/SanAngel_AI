"""Admin upload, review, and confirmation stay authenticated and branch-scoped."""

from datetime import date
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select

from backend.app.api.routes.admin_price_import import _service
from backend.app.core.admin_roles import AdminRole
from backend.app.db.models.branch import Branch
from backend.app.db.models.price import Price
from backend.app.db.models.price_import_audit import PriceImportAudit
from backend.app.db.models.product import Product
from backend.app.schemas.price_import_template import PRICE_IMPORT_COLUMNS
from backend.app.services.admin_auth_service import AdminAuthService
from backend.app.services.branch_scope import BranchScope
from backend.app.services.price_import_transaction import PriceImportTransactionService
from backend.tests.test_admin_commercial import PASSWORD, login, settings

pytest_plugins = ("backend.tests.test_admin_commercial",)

BASE = "/api/v1/admin/price-import"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook(product_id: int, *, branch: str = "sucursal-uno", name: str = "Aguja") -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Precios"
    sheet["A3"] = "schema_version"
    sheet["B3"] = 1
    sheet["D3"] = "branch_code"
    sheet["E3"] = branch
    for column, header in enumerate(PRICE_IMPORT_COLUMNS, start=1):
        sheet.cell(6, column, header)
    for column, value in enumerate((product_id, name, "kg", 123.45, date(2026, 1, 15)), 1):
        sheet.cell(7, column, value)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def upload(
    client: TestClient,
    content: bytes,
    csrf: dict[str, str] | None = None,
    filename: str = "precios.xlsx",
):
    return client.post(
        f"{BASE}/preview",
        content=content,
        headers={"Content-Type": XLSX, "X-File-Name": filename, **(csrf or {})},
    )


def test_authorization_csrf_branch_limit_and_one_time_confirmation(app, sessions) -> None:
    auth = AdminAuthService(sessions, settings=settings())
    auth.create_user("viewer", PASSWORD)
    auth.create_user("editor", PASSWORD, role=AdminRole.EDITOR)
    with sessions.begin() as session:
        own = session.scalar(select(Branch).where(Branch.code == "sucursal-uno"))
        foreign = session.scalar(select(Branch).where(Branch.code == "sucursal-dos"))
        assert own is not None and foreign is not None
        item = Product(branch_id=own.id, name="Aguja", category="Res")
        other = Product(branch_id=foreign.id, name="Ajeno", category="Res")
        session.add_all([item, other])
        session.flush()
        own_id, foreign_id = item.id, other.id
    app.dependency_overrides[_service] = lambda: PriceImportTransactionService(
        sessions, branch_scope=BranchScope("sucursal-uno")
    )
    try:
        with (
            TestClient(app, base_url="https://testserver") as anonymous,
            TestClient(app, base_url="https://testserver") as viewer,
            TestClient(app, base_url="https://testserver") as editor,
        ):
            content = workbook(own_id)
            assert upload(anonymous, content).status_code == 401
            assert upload(anonymous, b"x" * (2 * 1024 * 1024 + 1)).status_code == 401
            viewer_csrf = login(viewer, "viewer")
            assert upload(viewer, content, viewer_csrf).status_code == 403
            csrf = login(editor, "editor")
            assert upload(editor, content).status_code == 403
            assert upload(editor, b"x" * (2 * 1024 * 1024 + 1)).status_code == 403
            assert upload(editor, b"x" * (2 * 1024 * 1024 + 1), csrf).status_code == 413
            assert (
                upload(editor, content, csrf, filename="precios.xlsm").json()["issues"][0]["code"]
                == "invalid_filename"
            )
            assert (
                upload(
                    editor, workbook(foreign_id, branch="sucursal-dos", name="Ajeno"), csrf
                ).json()["issues"][0]["code"]
                == "branch_mismatch"
            )
            assert (
                upload(editor, workbook(foreign_id, name="Ajeno"), csrf).json()["issues"][0]["code"]
                == "product_unavailable"
            )
            preview_response = upload(editor, content, csrf)
            assert preview_response.status_code == 200
            assert preview_response.headers["cache-control"] == "no-store"
            preview = preview_response.json()
            assert preview["branch_code"] == "sucursal-uno"
            assert preview["summary"] == {"new": 1, "changed": 0, "unchanged": 0, "errors": 0}
            token = preview["preview_id"]
            assert token and len(token) == 32
            body = {"preview_id": token, "confirmed": True}
            assert editor.post(f"{BASE}/confirm", json=body).status_code == 403
            assert viewer.post(f"{BASE}/confirm", json=body, headers=viewer_csrf).status_code == 403
            assert (
                editor.post(
                    f"{BASE}/confirm", json={**body, "branch_code": "sucursal-dos"}, headers=csrf
                ).status_code
                == 422
            )
            receipt = editor.post(f"{BASE}/confirm", json=body, headers=csrf)
            assert receipt.status_code == 200
            assert receipt.json()["created"] == 1
            assert receipt.json()["branch_code"] == "sucursal-uno"
            assert editor.post(f"{BASE}/confirm", json=body, headers=csrf).status_code == 409
        with sessions() as session:
            assert (
                str(session.scalar(select(Price.amount).where(Price.product_id == own_id)))
                == "123.45"
            )
            assert session.scalar(select(Price).where(Price.product_id == foreign_id)) is None
            audits = session.scalars(select(PriceImportAudit)).all()
            assert len(audits) == 1 and audits[0].status == "success"
    finally:
        app.dependency_overrides.pop(_service, None)


def test_pending_preview_is_bound_to_session_and_revalidates_prices(app, sessions) -> None:
    auth = AdminAuthService(sessions, settings=settings())
    auth.create_user("editor", PASSWORD, role=AdminRole.EDITOR)
    with sessions.begin() as session:
        own = session.scalar(select(Branch).where(Branch.code == "sucursal-uno"))
        assert own is not None
        item = Product(branch_id=own.id, name="Aguja", category="Res")
        session.add(item)
        session.flush()
        item_id = item.id
    app.dependency_overrides[_service] = lambda: PriceImportTransactionService(
        sessions, branch_scope=BranchScope("sucursal-uno")
    )
    try:
        with (
            TestClient(app, base_url="https://testserver") as first,
            TestClient(app, base_url="https://testserver") as second,
        ):
            first_csrf = login(first, "editor")
            second_csrf = login(second, "editor")
            token = upload(first, workbook(item_id), first_csrf).json()["preview_id"]
            body = {"preview_id": token, "confirmed": True}
            assert second.post(f"{BASE}/confirm", json=body, headers=second_csrf).status_code == 409
            with sessions.begin() as session:
                own = session.scalar(select(Branch).where(Branch.code == "sucursal-uno"))
                session.add(Price(branch_id=own.id, product_id=item_id, unit="kg", amount="99.00"))
            assert first.post(f"{BASE}/confirm", json=body, headers=first_csrf).status_code == 409
        with sessions() as session:
            assert (
                str(session.scalar(select(Price.amount).where(Price.product_id == item_id)))
                == "99.00"
            )
            assert session.scalar(select(PriceImportAudit.status)) == "rejected"
    finally:
        app.dependency_overrides.pop(_service, None)
