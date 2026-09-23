"""Preview only the scoped price impact; never write while reviewing a file."""

import json
from collections.abc import Generator
from dataclasses import asdict
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from openpyxl import Workbook
from sqlalchemy import event, select
from sqlalchemy.orm import Session, sessionmaker

from backend.app.core.config import AssistantSettings, get_database_settings
from backend.app.db.models.price import Price
from backend.app.db.models.product import Product
from backend.app.db.session import create_database_engine, create_database_session_factory
from backend.app.repositories.branch_repository import BranchRepository
from backend.app.repositories.price_repository import PriceRepository
from backend.app.repositories.product_repository import ProductRepository
from backend.app.schemas.branch import BranchData
from backend.app.schemas.price import PriceData
from backend.app.schemas.price_import_template import PRICE_IMPORT_COLUMNS
from backend.app.schemas.product import ProductData
from backend.app.services.branch_scope import BranchScope
from backend.app.services.price_import_preview import PriceImportPreviewService

ROOT = Path(__file__).resolve().parents[2]
SCOPE = BranchScope.from_settings(AssistantSettings(assistant_branch_code="sucursal-demo"))
VERIFIED_ON = date(2026, 1, 15)


def workbook_bytes(
    rows: list[tuple[int, str, str, int | float, date]], *, branch: str = "sucursal-demo"
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Precios"
    sheet["A3"] = "schema_version"
    sheet["B3"] = 1
    sheet["D3"] = "branch_code"
    sheet["E3"] = branch
    for column, header in enumerate(PRICE_IMPORT_COLUMNS, start=1):
        sheet.cell(6, column, header)
    for source_row, values in enumerate(rows, start=7):
        for column, value in enumerate(values, start=1):
            sheet.cell(source_row, column, value)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def make_branch(session: Session, code: str):
    return BranchRepository(session).create(
        BranchData(
            code=code,
            name=f"Sucursal {code}",
            address="Dirección ficticia",
            phone=None,
            business_hours="Lunes a viernes",
        )
    )


@pytest.fixture
def catalog(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> Generator[tuple[sessionmaker[Session], dict[str, int]]]:
    database_url = f"sqlite+pysqlite:///{(tmp_path / 'preview.db').as_posix()}"
    monkeypatch.setenv("DATABASE_URL", database_url)
    get_database_settings.cache_clear()
    command.upgrade(Config(str(ROOT / "alembic.ini")), "head")
    engine = create_database_engine(database_url)
    factory = create_database_session_factory(engine)
    identifiers: dict[str, int] = {}
    try:
        with factory.begin() as session:
            branch = make_branch(session, "sucursal-demo")
            other_branch = make_branch(session, "otra-sucursal")
            products = ProductRepository(session, branch=branch)
            for name in ("Nuevo", "Cambio", "Igual", "Inactivo", "Nombre correcto"):
                product = products.create(ProductData(name=name, category="Res"))
                identifiers[name] = product.id
                if name == "Inactivo":
                    products.set_active(product, is_active=False)
                if name in {"Cambio", "Igual"}:
                    amount = "100.00" if name == "Cambio" else "30.00"
                    PriceRepository(session, branch=branch).create(
                        product, PriceData(amount=amount, unit="kg")
                    )
            foreign = ProductRepository(session, branch=other_branch).create(
                ProductData(name="Producto ajeno", category="Res")
            )
            identifiers["foreign"] = foreign.id
        yield factory, identifiers
    finally:
        engine.dispose()
        get_database_settings.cache_clear()


def price_snapshot(factory: sessionmaker[Session]) -> tuple[tuple[int, int, str, Decimal], ...]:
    with factory() as session:
        statement = select(Price.id, Price.product_id, Price.unit, Price.amount).order_by(Price.id)
        return tuple(session.execute(statement).all())


def test_preview_shows_new_changed_and_unchanged_without_persisting(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    before = price_snapshot(factory)
    content = workbook_bytes(
        [
            (identifiers["Nuevo"], "Nuevo", "kg", 75, VERIFIED_ON),
            (identifiers["Cambio"], "Cambio", "kg", 120, VERIFIED_ON),
            (identifiers["Igual"], "Igual", "kg", 30, VERIFIED_ON),
        ]
    )
    statements: list[str] = []

    def record_sql(
        _connection, _cursor, statement: str, _parameters, _context, _executemany
    ) -> None:
        statements.append(statement)

    engine = factory.kw["bind"]
    event.listen(engine, "before_cursor_execute", record_sql)
    try:
        with factory() as session:
            preview = PriceImportPreviewService(session, branch_scope=SCOPE).preview(
                content, filename="precios.xlsx"
            )
            assert not session.new and not session.dirty and not session.deleted
    finally:
        event.remove(engine, "before_cursor_execute", record_sql)

    assert preview.is_valid
    assert (
        preview.new_count,
        preview.changed_count,
        preview.unchanged_count,
        preview.error_count,
    ) == (
        1,
        1,
        1,
        0,
    )
    assert [item.action for item in preview.items] == ["new", "changed", "unchanged"]
    assert [(item.current_price_mxn, item.proposed_price_mxn) for item in preview.items] == [
        (None, Decimal("75")),
        (Decimal("100.00"), Decimal("120")),
        (Decimal("30.00"), Decimal("30")),
    ]
    assert all(item.verified_on == VERIFIED_ON for item in preview.items)
    assert statements and all(
        statement.lstrip().upper().startswith("SELECT") for statement in statements
    )
    review = preview.to_review_data()
    assert review["summary"] == {"new": 1, "changed": 1, "unchanged": 1, "errors": 0}
    assert review["items"][1]["current_price_mxn"] == "100.00"
    assert review["items"][1]["proposed_price_mxn"] == "120"
    json.dumps(review)
    assert price_snapshot(factory) == before


def test_preview_reports_scoped_product_errors_without_echoing_file_values(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    before = price_snapshot(factory)
    content = workbook_bytes(
        [
            (identifiers["Nuevo"], "Nuevo", "kg", 75, VERIFIED_ON),
            (identifiers["Nombre correcto"], "Nombre alterado", "kg", 20, VERIFIED_ON),
            (identifiers["Inactivo"], "Inactivo", "kg", 20, VERIFIED_ON),
            (identifiers["foreign"], "Producto ajeno", "kg", 20, VERIFIED_ON),
        ]
    )

    with factory() as session:
        preview = PriceImportPreviewService(session, branch_scope=SCOPE).preview(
            content, filename="precios.xlsx"
        )

    assert not preview.is_valid
    assert preview.new_count == 1
    assert preview.error_count == 3
    assert [(issue.row, issue.field, issue.code) for issue in preview.issues] == [
        (8, "product_name", "product_name_mismatch"),
        (9, "product_id", "product_unavailable"),
        (10, "product_id", "product_unavailable"),
    ]
    assert "Nombre alterado" not in repr(preview)
    assert "Producto ajeno" not in repr(preview)
    assert "Nombre alterado" not in json.dumps(preview.to_review_data())
    assert "Producto ajeno" not in json.dumps(preview.to_review_data())
    assert set(preview.to_review_data()) == {"is_valid", "summary", "items", "issues"}
    assert set(asdict(preview.items[0])) == {
        "source_row",
        "product_id",
        "product_name",
        "unit",
        "current_price_mxn",
        "proposed_price_mxn",
        "verified_on",
        "action",
    }
    assert price_snapshot(factory) == before


@pytest.mark.parametrize("branch,price", (("otra-sucursal", 75), ("sucursal-demo", -1)))
def test_rejected_workbook_does_not_query_or_mutate_database(
    catalog: tuple[sessionmaker[Session], dict[str, int]], branch: str, price: int
) -> None:
    factory, identifiers = catalog
    content = workbook_bytes(
        [(identifiers["Nuevo"], "Nuevo", "kg", price, VERIFIED_ON)], branch=branch
    )
    statements: list[str] = []

    def record_sql(
        _connection, _cursor, statement: str, _parameters, _context, _executemany
    ) -> None:
        statements.append(statement)

    engine = factory.kw["bind"]
    event.listen(engine, "before_cursor_execute", record_sql)
    try:
        with factory() as session:
            preview = PriceImportPreviewService(session, branch_scope=SCOPE).preview(
                content, filename="precios.xlsx"
            )
    finally:
        event.remove(engine, "before_cursor_execute", record_sql)

    assert preview.items == ()
    assert preview.error_count == 1
    assert statements == []


def test_unconfigured_branch_blocks_preview(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    content = workbook_bytes(
        [(identifiers["Nuevo"], "Nuevo", "kg", 75, VERIFIED_ON)], branch="sin-sucursal"
    )
    with factory() as session:
        preview = PriceImportPreviewService(
            session, branch_scope=BranchScope("sin-sucursal")
        ).preview(content, filename="precios.xlsx")

    assert preview.items == ()
    assert preview.issues[0].code == "branch_unavailable"


def test_preview_rejects_session_with_pending_changes(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    content = workbook_bytes([(identifiers["Nuevo"], "Nuevo", "kg", 75, VERIFIED_ON)])
    with factory() as session:
        session.add(Product(branch_id=1, name="Pendiente", category="Res", is_active=True))
        with pytest.raises(ValueError, match="clean database session"):
            PriceImportPreviewService(session, branch_scope=SCOPE).preview(
                content, filename="precios.xlsx"
            )
        session.rollback()
