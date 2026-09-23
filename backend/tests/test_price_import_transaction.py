"""Confirmed price imports are scoped, exact, and atomic."""

from collections.abc import Generator
from datetime import UTC, date
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from alembic import command
from alembic.config import Config
from openpyxl import Workbook
from sqlalchemy import event, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, sessionmaker

from backend.app.core.config import AssistantSettings, get_database_settings
from backend.app.db.models.price import Price
from backend.app.db.models.price_import_audit import PriceImportAudit
from backend.app.db.session import create_database_engine, create_database_session_factory
from backend.app.repositories.branch_repository import BranchRepository
from backend.app.repositories.price_import_audit_repository import (
    ImportActor,
    PriceImportAuditRepository,
)
from backend.app.repositories.price_repository import PriceRepository
from backend.app.repositories.product_repository import ProductRepository
from backend.app.schemas.branch import BranchData
from backend.app.schemas.price import PriceData
from backend.app.schemas.price_import_template import PRICE_IMPORT_COLUMNS
from backend.app.schemas.product import ProductData
from backend.app.services.branch_scope import BranchScope
from backend.app.services.price_import_parser import PriceImportIssue
from backend.app.services.price_import_transaction import (
    PriceImportAuditError,
    PriceImportConfirmationError,
    PriceImportStalePreviewError,
    PriceImportTransactionService,
    PriceImportValidationError,
    PriceImportWriteError,
)

ROOT = Path(__file__).resolve().parents[2]
SCOPE = BranchScope.from_settings(AssistantSettings(assistant_branch_code="sucursal-demo"))
ACTOR = ImportActor("staff-demo")
VERIFIED_ON = date(2026, 1, 15)


def workbook_bytes(
    rows: list[tuple[int, str, str, int | float, date]], *, branch: str = "sucursal-demo"
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Precios"
    sheet["A3"] = "schema_version"
    sheet["B3"] = 1
    sheet["D3"] = "branch_code"
    sheet["E3"] = branch
    for column, header in enumerate(PRICE_IMPORT_COLUMNS, start=1):
        sheet.cell(6, column, header)
    for source_row, values in enumerate(rows, start=7):
        for column, value in enumerate(values, start=1):
            sheet.cell(source_row, column, value)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def make_branch(session: Session, code: str):
    return BranchRepository(session).create(
        BranchData(
            code=code,
            name=f"Sucursal {code}",
            address="Dirección ficticia",
            phone=None,
            business_hours="Lunes a viernes",
        )
    )


@pytest.fixture
def catalog(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> Generator[tuple[sessionmaker[Session], dict[str, int]]]:
    database_url = f"sqlite+pysqlite:///{(tmp_path / 'import.db').as_posix()}"
    monkeypatch.setenv("DATABASE_URL", database_url)
    get_database_settings.cache_clear()
    command.upgrade(Config(str(ROOT / "alembic.ini")), "head")
    engine = create_database_engine(database_url)
    factory = create_database_session_factory(engine)
    identifiers: dict[str, int] = {}
    try:
        with factory.begin() as session:
            branch = make_branch(session, "sucursal-demo")
            other_branch = make_branch(session, "otra-sucursal")
            products = ProductRepository(session, branch=branch)
            for name in ("Nuevo", "Cambio", "Igual", "Cambiante"):
                product = products.create(ProductData(name=name, category="Res"))
                identifiers[name] = product.id
                if name in {"Cambio", "Igual"}:
                    amount = "100.00" if name == "Cambio" else "30.00"
                    PriceRepository(session, branch=branch).create(
                        product, PriceData(amount=amount, unit="kg")
                    )
            foreign = ProductRepository(session, branch=other_branch).create(
                ProductData(name="Ajeno", category="Res")
            )
            identifiers["foreign"] = foreign.id
            PriceRepository(session, branch=other_branch).create(
                foreign, PriceData(amount="999.00", unit="kg")
            )
        yield factory, identifiers
    finally:
        engine.dispose()
        get_database_settings.cache_clear()


def prices(factory: sessionmaker[Session]) -> tuple[tuple[int, str, Decimal], ...]:
    with factory() as session:
        return tuple(
            session.execute(
                select(Price.product_id, Price.unit, Price.amount).order_by(Price.product_id)
            ).all()
        )


def import_content(identifiers: dict[str, int]) -> bytes:
    return workbook_bytes(
        [
            (identifiers["Nuevo"], "Nuevo", "kg", 75.25, VERIFIED_ON),
            (identifiers["Cambio"], "Cambio", "kg", 120.55, VERIFIED_ON),
            (identifiers["Igual"], "Igual", "kg", 30, VERIFIED_ON),
        ]
    )


def test_confirmed_import_upserts_exact_prices_and_returns_receipt(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    service = PriceImportTransactionService(factory, branch_scope=SCOPE)
    content = import_content(identifiers)
    before = prices(factory)
    prepared = service.prepare(content, filename="precios.xlsx")

    assert prepared.preview.is_valid
    assert prices(factory) == before
    receipt = service.confirm(
        content, filename="precios.xlsx", prepared=prepared, confirmed=True, actor=ACTOR
    )

    assert (receipt.created, receipt.updated, receipt.unchanged) == (1, 1, 1)
    assert receipt.file_sha256 == sha256(content).hexdigest()
    assert receipt.branch_code == SCOPE.branch_code
    report = service.get_report(receipt.audit_id)
    assert report is not None
    assert report.attempt_id == receipt.attempt_id
    assert report.actor_id == ACTOR.actor_id
    assert report.occurred_at.tzinfo == UTC
    assert report.status == "success"
    assert (report.created, report.updated, report.unchanged, report.error_count) == (1, 1, 1, 0)
    assert report.file_sha256 == receipt.file_sha256
    assert report.errors == ()
    assert "precios.xlsx" not in str(report.to_review_data())
    with factory() as session:
        stored = session.scalar(
            select(PriceImportAudit).where(PriceImportAudit.id == receipt.audit_id)
        )
        assert stored is not None
        assert stored.errors_json == "[]"
    assert prices(factory) == (
        (identifiers["Nuevo"], "kg", Decimal("75.25")),
        (identifiers["Cambio"], "kg", Decimal("120.55")),
        (identifiers["Igual"], "kg", Decimal("30.00")),
        (identifiers["foreign"], "kg", Decimal("999.00")),
    )

    with pytest.raises(PriceImportStalePreviewError):
        service.confirm(
            content, filename="precios.xlsx", prepared=prepared, confirmed=True, actor=ACTOR
        )
    refreshed = service.prepare(content, filename="precios.xlsx")
    repeat = service.confirm(
        content, filename="precios.xlsx", prepared=refreshed, confirmed=True, actor=ACTOR
    )
    assert (repeat.created, repeat.updated, repeat.unchanged) == (0, 0, 3)
    assert len(service.list_reports()) == 3


@pytest.mark.parametrize("confirmation", (False, None, "true", 1))
def test_explicit_boolean_confirmation_is_required(
    catalog: tuple[sessionmaker[Session], dict[str, int]], confirmation: object
) -> None:
    factory, identifiers = catalog
    service = PriceImportTransactionService(factory, branch_scope=SCOPE)
    content = import_content(identifiers)
    prepared = service.prepare(content, filename="precios.xlsx")
    before = prices(factory)

    with pytest.raises(PriceImportConfirmationError):
        service.confirm(
            content, filename="precios.xlsx", prepared=prepared, confirmed=confirmation, actor=ACTOR
        )
    assert prices(factory) == before
    report = PriceImportTransactionService(factory, branch_scope=SCOPE).list_reports()[0]
    assert report.status == "rejected"
    assert report.error_code == "confirmation_rejected"


def test_confirmation_is_bound_to_file_name_bytes_and_branch(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    service = PriceImportTransactionService(factory, branch_scope=SCOPE)
    content = import_content(identifiers)
    prepared = service.prepare(content, filename="precios.xlsx")
    before = prices(factory)
    other_content = workbook_bytes([(identifiers["Nuevo"], "Nuevo", "kg", 70, VERIFIED_ON)])

    for candidate_content, filename, scope in (
        (other_content, "precios.xlsx", SCOPE),
        (content, "otro.xlsx", SCOPE),
        (content, "precios.xlsx", BranchScope("otra-sucursal")),
    ):
        with pytest.raises(PriceImportConfirmationError):
            PriceImportTransactionService(factory, branch_scope=scope).confirm(
                candidate_content, filename=filename, prepared=prepared, confirmed=True, actor=ACTOR
            )
    assert prices(factory) == before


def test_invalid_or_foreign_product_never_writes(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    service = PriceImportTransactionService(factory, branch_scope=SCOPE)
    before = prices(factory)
    for content in (
        workbook_bytes([(identifiers["Nuevo"], "Nuevo", "kg", -1, VERIFIED_ON)]),
        workbook_bytes([(identifiers["foreign"], "Ajeno", "kg", 10, VERIFIED_ON)]),
        workbook_bytes(
            [(identifiers["Nuevo"], "Nuevo", "kg", 10, VERIFIED_ON)],
            branch="otra-sucursal",
        ),
    ):
        prepared = service.prepare(content, filename="precios.xlsx")
        assert not prepared.preview.is_valid
        assert prepared.file_sha256 == ""
        with pytest.raises(PriceImportValidationError) as captured:
            service.confirm(
                content, filename="precios.xlsx", prepared=prepared, confirmed=True, actor=ACTOR
            )
        assert captured.value.issues
        report = service.list_reports()[0]
        assert report.status == "rejected"
        assert report.error_count >= 1
        assert report.file_sha256 == sha256(content).hexdigest()
    assert prices(factory) == before


def test_malicious_package_is_rejected_without_price_writes(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    service = PriceImportTransactionService(factory, branch_scope=SCOPE)
    before = prices(factory)
    output = BytesIO()
    with ZipFile(BytesIO(import_content(identifiers))) as source, ZipFile(output, "w") as target:
        for member in source.infolist():
            target.writestr(member, source.read(member.filename))
        target.writestr("xl/vbaProject.bin", b"untrusted workbook payload")
    content = output.getvalue()

    prepared = service.prepare(content, filename="precios.xlsx")
    assert prepared.preview.issues[0].code == "unsafe_package"
    with pytest.raises(PriceImportValidationError):
        service.confirm(
            content, filename="precios.xlsx", prepared=prepared, confirmed=True, actor=ACTOR
        )

    assert prices(factory) == before
    report = service.list_reports()[0]
    assert report.status == "rejected"
    assert report.errors[0].code == "unsafe_package"
    assert "untrusted workbook payload" not in str(report.to_review_data())


def test_changed_catalog_or_price_blocks_stale_preview(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    service = PriceImportTransactionService(factory, branch_scope=SCOPE)
    content = import_content(identifiers)
    prepared = service.prepare(content, filename="precios.xlsx")
    with factory.begin() as session:
        branch = BranchRepository(session).get_by_code(SCOPE.branch_code)
        product = ProductRepository(session, branch=branch).get_by_id(identifiers["Cambio"])
        current = PriceRepository(session, branch=branch).get_for_product(product, unit="kg")
        PriceRepository(session, branch=branch).update(
            current, PriceData(amount="105.00", unit="kg")
        )
    before = prices(factory)

    with pytest.raises(PriceImportStalePreviewError):
        service.confirm(
            content, filename="precios.xlsx", prepared=prepared, confirmed=True, actor=ACTOR
        )
    assert prices(factory) == before


def test_product_renamed_after_preview_blocks_every_write(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    service = PriceImportTransactionService(factory, branch_scope=SCOPE)
    content = import_content(identifiers)
    prepared = service.prepare(content, filename="precios.xlsx")
    with factory.begin() as session:
        branch = BranchRepository(session).get_by_code(SCOPE.branch_code)
        product = ProductRepository(session, branch=branch).get_by_id(identifiers["Cambio"])
        ProductRepository(session, branch=branch).update(
            product, ProductData(name="Renombrado", category="Res")
        )
    before = prices(factory)

    with pytest.raises(PriceImportValidationError) as captured:
        service.confirm(
            content, filename="precios.xlsx", prepared=prepared, confirmed=True, actor=ACTOR
        )
    assert captured.value.issues[0].code == "product_name_mismatch"
    assert prices(factory) == before


def test_late_database_failure_rolls_back_earlier_price_write(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    service = PriceImportTransactionService(factory, branch_scope=SCOPE)
    content = import_content(identifiers)
    prepared = service.prepare(content, filename="precios.xlsx")
    before = prices(factory)
    writes = 0

    def fail_second_write(_connection, _cursor, statement, parameters, _context, _executemany):
        nonlocal writes
        if statement.lstrip().upper().startswith(("INSERT INTO PRICES", "UPDATE PRICES")):
            writes += 1
            if writes == 2:
                raise IntegrityError(statement, parameters, Exception("synthetic failure"))

    engine = factory.kw["bind"]
    event.listen(engine, "before_cursor_execute", fail_second_write)
    try:
        with pytest.raises(PriceImportWriteError, match="could not be committed"):
            service.confirm(
                content, filename="precios.xlsx", prepared=prepared, confirmed=True, actor=ACTOR
            )
    finally:
        event.remove(engine, "before_cursor_execute", fail_second_write)

    assert writes == 2
    assert prices(factory) == before
    report = service.list_reports()[0]
    assert report.status == "failed"
    assert report.error_code == "database_error"
    assert (report.created, report.updated, report.unchanged) == (0, 0, 0)


def test_audit_insert_failure_rolls_back_prices_and_records_failed_attempt(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    service = PriceImportTransactionService(factory, branch_scope=SCOPE)
    content = import_content(identifiers)
    prepared = service.prepare(content, filename="precios.xlsx")
    before = prices(factory)
    audit_writes = 0

    def fail_first_audit(_connection, _cursor, statement, parameters, _context, _executemany):
        nonlocal audit_writes
        if statement.lstrip().upper().startswith("INSERT INTO PRICE_IMPORT_AUDITS"):
            audit_writes += 1
            if audit_writes == 1:
                raise IntegrityError(statement, parameters, Exception("synthetic audit failure"))

    engine = factory.kw["bind"]
    event.listen(engine, "before_cursor_execute", fail_first_audit)
    try:
        with pytest.raises(PriceImportWriteError):
            service.confirm(
                content, filename="precios.xlsx", prepared=prepared, confirmed=True, actor=ACTOR
            )
    finally:
        event.remove(engine, "before_cursor_execute", fail_first_audit)

    assert audit_writes == 2
    assert prices(factory) == before
    assert [report.status for report in service.list_reports()] == ["failed"]


def test_unavailable_audit_store_blocks_import_with_safe_error(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    service = PriceImportTransactionService(factory, branch_scope=SCOPE)
    content = import_content(identifiers)
    prepared = service.prepare(content, filename="precios.xlsx")
    before = prices(factory)

    def fail_audit(_connection, _cursor, statement, parameters, _context, _executemany):
        if statement.lstrip().upper().startswith("INSERT INTO PRICE_IMPORT_AUDITS"):
            raise IntegrityError(statement, parameters, Exception("synthetic secret detail"))

    engine = factory.kw["bind"]
    event.listen(engine, "before_cursor_execute", fail_audit)
    try:
        with pytest.raises(PriceImportAuditError, match="could not be audited") as captured:
            service.confirm(
                content, filename="precios.xlsx", prepared=prepared, confirmed=True, actor=ACTOR
            )
    finally:
        event.remove(engine, "before_cursor_execute", fail_audit)

    assert "synthetic secret detail" not in str(captured.value)
    assert prices(factory) == before
    assert service.list_reports() == ()


def test_audit_reports_are_scoped_and_exclude_workbook_content(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, identifiers = catalog
    service = PriceImportTransactionService(factory, branch_scope=SCOPE)
    content = import_content(identifiers)
    prepared = service.prepare(content, filename="precios.xlsx")
    receipt = service.confirm(
        content, filename="precios.xlsx", prepared=prepared, confirmed=True, actor=ACTOR
    )

    other = PriceImportTransactionService(factory, branch_scope=BranchScope("otra-sucursal"))
    assert other.get_report(receipt.audit_id) is None
    assert other.list_reports() == ()
    report_data = service.get_report(receipt.audit_id).to_review_data()
    assert set(report_data) == {
        "audit_id",
        "attempt_id",
        "branch_code",
        "actor_id",
        "occurred_at",
        "file_sha256",
        "status",
        "summary",
        "error_code",
        "errors_truncated",
        "errors",
    }
    assert "precios.xlsx" not in str(report_data)
    assert "Nuevo" not in str(report_data)
    assert str(content[:20]) not in str(report_data)


def test_audit_limits_stored_row_errors_without_losing_total_count(
    catalog: tuple[sessionmaker[Session], dict[str, int]],
) -> None:
    factory, _ = catalog
    issues = tuple(PriceImportIssue(row, "price_mxn", "out_of_range") for row in range(1, 206))
    with factory.begin() as session:
        audit = PriceImportAuditRepository(session, branch_scope=SCOPE).record(
            attempt_id="a" * 32,
            actor=ACTOR,
            file_sha256=None,
            status="rejected",
            error_code="validation_failed",
            issues=issues,
        )
        audit_id = audit.id
    report = PriceImportTransactionService(factory, branch_scope=SCOPE).get_report(audit_id)

    assert report.error_count == 205
    assert len(report.errors) == 200
    assert report.to_review_data()["errors_truncated"] is True


@pytest.mark.parametrize("actor_id", ("", "user@example.com", "123456789", "x" * 65))
def test_audit_actor_requires_opaque_backend_identifier(actor_id: str) -> None:
    with pytest.raises(ValueError, match="opaque staff identifier"):
        ImportActor(actor_id)
